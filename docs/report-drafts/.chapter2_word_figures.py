"""Insert verified Chapter 2 figures into the existing formatted Word draft."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "Chapter_2_E-Filing_Draft.docx"
FIGURES = ROOT / ".chapter2-figures"

PLACEHOLDERS = [
    ("pdca-unified", 1, "วงจร PDCA และงานหลักในระบบ E-Filing เดียว"),
    ("dar-ack", 3, "วงจร DAR และการรับทราบเอกสารตามฉบับ"),
    ("audit-finding", 4, "ความสัมพันธ์ระหว่างการตรวจติดตาม เกณฑ์ หลักฐาน และ Finding"),
    ("car-pending-closure", 5, "สถานะ CAR/NCR ในระบบทดลองถึงขั้นรอปิดใบ"),
    ("architecture-unified", 7, "บริบทและส่วนประกอบหลักของระบบ E-Filing เดียว"),
    ("data-relations-unified", 8, "ความสัมพันธ์ของข้อมูลหลักระดับแนวคิด"),
    ("dashboard-unified", 9, "การสรุปข้อมูลแดชบอร์ดจากหกหมวดงาน"),
    ("traceability-unified", 10, "เส้นทางการตรวจสอบย้อนกลับของข้อมูล"),
]


def paragraph_after(paragraph: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    return Paragraph(element, paragraph._parent)


def image_paragraph(paragraph: Paragraph, slug: str) -> None:
    paragraph.clear()
    paragraph.style = "E-Filing Thai Body"
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = 0
    fmt.keep_with_next = True
    fmt.keep_together = True
    paragraph.add_run().add_picture(str(FIGURES / f"{slug}.png"), width=Inches(5.6))


def caption_after(paragraph: Paragraph, number: int, title: str) -> None:
    caption = paragraph_after(paragraph)
    caption.style = "E-Filing Thai Body"
    fmt = caption.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = 0
    fmt.keep_together = True
    fmt.keep_with_next = False
    caption.add_run(f"รูปที่ 2.{number}").bold = True
    caption.add_run(f" {title}")


def main() -> None:
    doc = Document(DOCX)
    assert len(doc.inline_shapes) == 1, "Expected only the already reused hierarchy image"
    old = [p for p in doc.paragraphs if p.text.startswith("[ต้องจัดทำรูปใหม่:")]
    assert len(old) == len(PLACEHOLDERS), f"Expected 8 figure placeholders, found {len(old)}"

    for paragraph, (slug, number, title) in zip(old, PLACEHOLDERS):
        image_paragraph(paragraph, slug)
        caption_after(paragraph, number, title)

    hierarchy = [p for p in doc.paragraphs if p.text.startswith("รูปที่ 2.1 ลำดับชั้นเอกสาร")]
    assert len(hierarchy) == 1, "Expected existing hierarchy figure caption"
    hierarchy[0].runs[0].text = hierarchy[0].runs[0].text.replace("รูปที่ 2.1", "รูปที่ 2.2", 1)

    rbac_anchor = [p for p in doc.paragraphs if p.text.startswith("ระบบ E-Filing ใช้บัญชีและบทบาทร่วมกัน")]
    assert len(rbac_anchor) == 1, "Expected RBAC prose paragraph"
    rbac_image = paragraph_after(rbac_anchor[0])
    image_paragraph(rbac_image, "rbac-unified")
    caption_after(rbac_image, 6, "หลักการควบคุมสิทธิ์ตามบทบาทและการแยกหน้าที่")

    doc.save(DOCX)
    print(f"Updated {DOCX} with {len(doc.inline_shapes)} total figures")


if __name__ == "__main__":
    main()
