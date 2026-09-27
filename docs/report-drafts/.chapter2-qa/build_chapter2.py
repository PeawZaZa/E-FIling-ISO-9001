"""Build Chapter 2 from the retained Group B Word formatting reference."""

import copy
import re
import sys
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import generate_report_docx_from_reference as base

W = base.W


def blocks(path):
    result = []
    rows = []
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            rows.append(cells)
            continue
        if rows:
            result.append(("table", rows))
            rows = []
        if line.startswith("# "):
            result.append(("chapter", line[2:]))
        elif line.startswith("## "):
            result.append(("major", line[3:]))
        elif line.startswith("### "):
            result.append(("sub", line[4:]))
        elif line.startswith("**ตารางที่ "):
            result.append(("table_caption", line))
        elif line.startswith("[ต้องจัดทำรูปใหม่:"):
            result.append(("figure_placeholder", line))
        elif re.match(r"^\[\d+\] ", line):
            result.append(("reference", line))
        else:
            result.append(("body", line))
    if rows:
        result.append(("table", rows))
    return result


def decorate(rpr, *, bold=False, italic=False, size=None):
    if rpr is None:
        rpr = OxmlElement("w:rPr")
    fonts = base.child(rpr, "rFonts")
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(W + name, "TH Sarabun New")
    if size:
        base.font_size(rpr, size)
    if bold:
        base.child(rpr, "b")
        base.child(rpr, "bCs")
    if italic:
        base.child(rpr, "i")
        base.child(rpr, "iCs")
    return rpr


def append_markdown(p, text, thai, latin):
    for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not token:
            continue
        is_bold = token.startswith("**") and token.endswith("**")
        is_italic = not is_bold and token.startswith("*") and token.endswith("*")
        if is_bold:
            token = token[2:-2]
        elif is_italic:
            token = token[1:-1]
        tr = decorate(copy.deepcopy(thai), bold=is_bold, italic=is_italic)
        lr = decorate(copy.deepcopy(latin), bold=is_bold, italic=is_italic)
        base.append_runs(p, token, tr, lr)


def add_style(doc, name, source_p, source_rpr, *, bold=False, indent=None,
              alignment=None, size=16):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    ppr = style._element.get_or_add_pPr()
    source = source_p.pPr
    if source is not None:
        for item in source:
            if item.tag != W + "rPr":
                ppr.append(copy.deepcopy(item))
    if indent is not None:
        base.child(ppr, "ind").set(W + "left", str(indent))
        base.child(ppr, "ind").set(W + "firstLine", "0")
    if alignment is not None:
        base.child(ppr, "jc").set(W + "val", alignment)
    rpr = style._element.get_or_add_rPr()
    for item in list(rpr):
        rpr.remove(item)
    for item in decorate(copy.deepcopy(source_rpr), bold=bold, size=size * 2):
        rpr.append(item)
    return style


def pstyle(p, style_id):
    ppr = base.child(p, "pPr")
    style = base.child(ppr, "pStyle")
    style.set(W + "val", style_id)
    return ppr


def make_paragraph(pattern, style, text, thai, latin, *, keep=False,
                   left=None, first=None):
    p = base.clean_clone(pattern)
    ppr = pstyle(p, style.style_id)
    if keep:
        base.child(ppr, "keepNext")
        base.child(ppr, "keepLines")
    if left is not None or first is not None:
        ind = base.child(ppr, "ind")
        if left is not None:
            ind.set(W + "left", str(left))
        if first is not None:
            ind.set(W + "firstLine", str(first))
    append_markdown(p, text, thai, latin)
    return p


def set_cell_text(p, text, style, thai, latin):
    base.replace_text(p, "", thai, latin)
    pstyle(p, style.style_id)
    append_markdown(p, text, thai, latin)


def make_table(source_table, values, style, thai, latin):
    if not values or len({len(row) for row in values}) != 1:
        raise ValueError("Malformed Markdown table")
    table = copy.deepcopy(source_table)
    ncol = len(values[0])
    if ncol != 3:
        raise ValueError("Unexpected Chapter 2 table width")
    widths = (1920, 3000, 3360)
    tblpr = table.find(W + "tblPr")
    tblw = base.child(tblpr, "tblW")
    tblw.set(W + "w", str(sum(widths)))
    tblw.set(W + "type", "dxa")
    grid = table.find(W + "tblGrid")
    for item in list(grid):
        grid.remove(item)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(W + "w", str(width))
        grid.append(col)
    row_pattern = table.find(W + "tr")
    for row in table.findall(W + "tr"):
        table.remove(row)
    for row_index, data in enumerate(values):
        row = copy.deepcopy(row_pattern)
        for height in row.findall(".//" + W + "trHeight"):
            height.getparent().remove(height)
        cells = row.findall(W + "tc")
        for extra in cells[ncol:]:
            row.remove(extra)
        for col_index, cell in enumerate(cells[:ncol]):
            tcpr = base.child(cell, "tcPr")
            tcw = base.child(tcpr, "tcW")
            tcw.set(W + "w", str(widths[col_index]))
            tcw.set(W + "type", "dxa")
            for shade in tcpr.findall(W + "shd"):
                tcpr.remove(shade)
            paras = cell.findall(W + "p")
            p = paras[0]
            for extra in paras[1:]:
                cell.remove(extra)
            tr = decorate(copy.deepcopy(thai), bold=row_index == 0, size=24)
            lr = decorate(copy.deepcopy(latin), bold=row_index == 0, size=24)
            set_cell_text(p, data[col_index], style, tr, lr)
            ppr = base.child(p, "pPr")
            base.child(ppr, "jc").set(W + "val", "center" if row_index == 0 else "left")
            ind = base.child(ppr, "ind")
            ind.set(W + "firstLine", "0")
            ind.set(W + "left", "0")
        if row_index == 0:
            base.child(base.child(row, "trPr"), "tblHeader")
        table.append(row)
    return table


def repack(intermediate, output, reference):
    with zipfile.ZipFile(intermediate) as source, zipfile.ZipFile(reference) as ref, \
            zipfile.ZipFile(output, "w") as dest:
        for info in source.infolist():
            name = info.filename
            if name.startswith("customXml/") or name == "docProps/custom.xml":
                continue
            data = source.read(name)
            if name in ("word/settings.xml", "word/theme/theme1.xml", "word/fontTable.xml"):
                data = ref.read(name)
            elif name in ("word/_rels/document.xml.rels", "_rels/.rels"):
                root = etree.fromstring(data)
                for rel in list(root):
                    target = rel.get("Target", "")
                    if "customXml/" in target or target.endswith("docProps/custom.xml"):
                        root.remove(rel)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name == "[Content_Types].xml":
                root = etree.fromstring(data)
                for part in list(root):
                    path = part.get("PartName", "")
                    if path.startswith("/customXml/") or path == "/docProps/custom.xml":
                        root.remove(part)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dest.writestr(info, data)


def build(reference, source, output):
    doc = Document(reference)
    ref = list(doc.paragraphs)
    patterns = {"chapter": ref[0]._p, "major": ref[1]._p, "body": ref[2]._p,
                "sub": ref[22]._p, "caption": ref[46]._p}
    table_pattern = doc.tables[0]._tbl
    body_thai = base.run_props(ref[2], 0)
    body_latin = base.run_props(ref[2], 1)
    main_thai = base.run_props(ref[1], 1)
    main_latin = base.run_props(ref[1], 0)
    chapter_thai = base.run_props(ref[0], 0)
    chapter_latin = base.run_props(ref[0], 1)
    styles = {
        "chapter": add_style(doc, "E-Filing Chapter Title", patterns["chapter"], chapter_thai, bold=True, alignment="center", size=18),
        "major": add_style(doc, "E-Filing Main Heading", patterns["major"], main_thai, bold=True, alignment="left"),
        "sub": add_style(doc, "E-Filing Subheading", patterns["sub"], main_thai, bold=True, indent=720, alignment="left"),
        "body": add_style(doc, "E-Filing Thai Body", patterns["body"], body_thai, alignment="thaiDistribute"),
        "table_caption": add_style(doc, "E-Filing Table Caption", patterns["caption"], body_thai, alignment="left"),
        "figure_placeholder": add_style(doc, "E-Filing Figure Placeholder", patterns["body"], body_thai, alignment="left"),
        "reference": add_style(doc, "E-Filing Reference Entry", patterns["body"], body_thai, alignment="left", size=15),
        "table": add_style(doc, "E-Filing Table Text", patterns["body"], body_thai, alignment="left", size=12),
    }
    body = doc._body._body
    sectpr = body.find(W + "sectPr")
    for item in list(body):
        if item is not sectpr:
            body.remove(item)
    after_table = False
    for kind, value in blocks(source):
        if kind == "table":
            item = make_table(table_pattern, value, styles["table"], body_thai, body_latin)
            after_table = True
        else:
            pattern = patterns.get(kind, patterns["body"])
            if kind == "chapter":
                thai, latin = chapter_thai, chapter_latin
            elif kind in ("major", "sub"):
                thai, latin = main_thai, main_latin
            else:
                thai, latin = body_thai, body_latin
            if kind == "reference":
                thai = decorate(copy.deepcopy(thai), size=30)
                latin = decorate(copy.deepcopy(latin), size=30)
            if kind == "table_caption":
                pattern = patterns["caption"]
            item = make_paragraph(pattern, styles[kind], value, thai, latin,
                                  keep=kind in ("chapter", "major", "sub", "table_caption"),
                                  left=720 if kind == "sub" else 0 if kind in ("figure_placeholder", "reference", "table_caption") else None,
                                  first=0 if kind in ("sub", "figure_placeholder", "reference", "table_caption") else None)
            ppr = base.child(item, "pPr")
            if kind == "body" and after_table:
                base.child(ppr, "spacing").set(W + "before", "120")
            after_table = False
            if kind in ("major", "sub", "table_caption", "figure_placeholder", "reference"):
                base.child(ppr, "jc").set(W + "val", "left")
            if kind == "table_caption":
                number = re.match(r"\*\*(ตารางที่ 2\.\d+)\*\*", value)
                if not number:
                    raise ValueError("Unexpected caption")
                for run in item.findall(W + "r"):
                    text = "".join(t.text or "" for t in run.findall(W + "t"))
                    if "ตารางที่" in text or re.match(r"2\.\d+", text):
                        base.child(base.child(run, "rPr"), "b")
                        base.child(base.child(run, "rPr"), "bCs")
        body.insert(len(body) - 1, item)
    base.page_numbers(doc)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "E-Filing ISO 9001 Chapter 2 Draft"
    doc.core_properties.subject = ""
    doc.core_properties.keywords = ""
    doc.core_properties.comments = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(dir=output.parent) as temp:
        intermediate = Path(temp) / "chapter2_working.docx"
        doc.save(intermediate)
        repack(intermediate, output, reference)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("Usage: build_chapter2.py REFERENCE.docx SOURCE.md OUTPUT.docx")
    build(*(Path(arg) for arg in sys.argv[1:]))
