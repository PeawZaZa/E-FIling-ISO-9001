"""Replace a Word-formatted chapter reference with report Markdown content.

The reference DOCX supplies layout/Word behavior only. No reference project
text is reused. Usage:

    python generate_report_docx_from_reference.py REFERENCE.docx SOURCE.md OUTPUT.docx
"""

import copy
import re
import sys
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm
from lxml import etree


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
THAI_OR_OTHER = re.compile(r"[\u0E00-\u0E7F]+|[^\u0E00-\u0E7F]+")


def child(parent, local_name):
    element = parent.find(W + local_name)
    if element is None:
        element = OxmlElement("w:" + local_name)
        parent.append(element)
    return element


def run_props(paragraph, index):
    return copy.deepcopy(paragraph.runs[index]._r.rPr)


def font_size(rpr, half_points):
    child(rpr, "sz").set(W + "val", str(half_points))
    child(rpr, "szCs").set(W + "val", str(half_points))


def clean_clone(paragraph_xml):
    result = copy.deepcopy(paragraph_xml)
    for attr in list(result.attrib):
        # Identifiers belong to the reference paragraph, not its clones.
        if etree.QName(attr).localname in ("paraId", "textId") or \
                etree.QName(attr).localname.startswith("rsid"):
            del result.attrib[attr]
    for element in list(result):
        if element.tag != W + "pPr":
            result.remove(element)
    return result


def append_runs(paragraph_xml, text, thai_props, latin_props):
    """Preserve text exactly while using Group B's script-aware run pattern."""
    for match in THAI_OR_OTHER.finditer(text):
        segment = match.group()
        is_thai = "\u0e00" <= segment[0] <= "\u0e7f"
        run = OxmlElement("w:r")
        run.append(copy.deepcopy(thai_props if is_thai else latin_props))
        t = OxmlElement("w:t")
        if segment[:1].isspace() or segment[-1:].isspace():
            t.set(XML_SPACE, "preserve")
        t.text = segment
        run.append(t)
        paragraph_xml.append(run)


def replace_text(paragraph_xml, text, thai_props, latin_props):
    for element in list(paragraph_xml):
        if element.tag != W + "pPr":
            paragraph_xml.remove(element)
    append_runs(paragraph_xml, text, thai_props, latin_props)


def parse_markdown(path):
    blocks = []
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        if line.startswith("| "):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r"-+", c) for c in cells):
                continue
            rows.append(cells)
            continue
        if rows:
            blocks.append(("table", rows))
            rows = []
        if line.startswith("# "):
            blocks.append(("chapter", line[2:]))
        elif line.startswith("## "):
            blocks.append(("major", line[3:]))
        elif line.startswith("### "):
            blocks.append(("sub", line[4:]))
        elif line.startswith("**ตารางที่ "):
            match = re.fullmatch(r"\*\*(.*?)\*\*(.*)", line)
            if match is None:
                raise ValueError("Unexpected table caption syntax")
            blocks.append(("caption", (match.group(1), match.group(2))))
        elif re.match(r"^1\.2\.\d+ ", line):
            blocks.append(("numbered", line))
        else:
            blocks.append(("body", line))
    if rows:
        blocks.append(("table", rows))
    return blocks


def make_paragraph(pattern, text, thai_props, latin_props, *, left=False,
                   keep_next=False):
    p = clean_clone(pattern)
    ppr = child(p, "pPr")
    if left:
        child(ppr, "jc").set(W + "val", "left")
    if keep_next:
        child(ppr, "keepNext")
    append_runs(p, text, thai_props, latin_props)
    return p


def make_table(pattern, values, normal_thai, normal_latin):
    if len(values) != 7 or any(len(row) != 3 for row in values):
        raise ValueError("Chapter 1 source schedule must have 7 rows and 3 columns")
    table = copy.deepcopy(pattern)
    widths = (720, 4800, 2760)  # fits inside the official A4 text frame
    tbl_pr = table.find(W + "tblPr")
    tbl_w = child(tbl_pr, "tblW")
    tbl_w.set(W + "w", str(sum(widths)))
    tbl_w.set(W + "type", "dxa")
    grid = table.find(W + "tblGrid")
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(W + "w", str(width))
        grid.append(col)
    rows = table.findall(W + "tr")
    if len(rows) != 7:
        raise ValueError("Reference schedule no longer has 7 rows")
    for row_num, row in enumerate(rows):
        cells = row.findall(W + "tc")
        for excess in cells[3:]:
            row.remove(excess)
        for col_num, cell in enumerate(cells[:3]):
            tc_pr = child(cell, "tcPr")
            tc_w = child(tc_pr, "tcW")
            tc_w.set(W + "w", str(widths[col_num]))
            tc_w.set(W + "type", "dxa")
            for shade in tc_pr.findall(W + "shd"):
                tc_pr.remove(shade)  # Group B's filled timeline cell is data-specific.
            paragraphs = cell.findall(W + "p")
            if not paragraphs:
                p = OxmlElement("w:p")
                cell.append(p)
            else:
                p = paragraphs[0]
                for excess in paragraphs[1:]:
                    cell.remove(excess)
            thai = copy.deepcopy(normal_thai)
            latin = copy.deepcopy(normal_latin)
            font_size(thai, 24)  # official table-content rule: 12 pt
            font_size(latin, 24)
            replace_text(p, values[row_num][col_num], thai, latin)
    return table


def page_numbers(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.header_distance = Inches(1)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = 2  # right
    header.paragraph_format.first_line_indent = Inches(0)
    run = header.add_run()
    rpr = run._r.get_or_add_rPr()
    fonts = child(rpr, "rFonts")
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(W + key, "TH Sarabun New")
    font_size(rpr, 32)
    field = OxmlElement("w:fldSimple")
    field.set(W + "instr", "PAGE")
    header._p.remove(run._r)
    field.append(run._r)
    cached = OxmlElement("w:t")
    cached.text = "1"
    run._r.append(cached)
    header._p.append(field)


def remove_unused_metadata(source, destination, reference):
    """Keep reference formatting parts but omit project metadata/custom XML."""
    exact_format_parts = ("word/styles.xml", "word/settings.xml",
                          "word/theme/theme1.xml", "word/fontTable.xml")
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(reference) as base, \
            zipfile.ZipFile(destination, "w") as result:
        for info in original.infolist():
            name = info.filename
            if name.startswith("customXml/") or name == "docProps/custom.xml":
                continue
            data = original.read(name)
            if name in exact_format_parts:
                # python-docx can reserialize these parts despite no style edit.
                # Restore the exact known-good Word formatting bytes.
                data = base.read(name)
            elif name in ("word/_rels/document.xml.rels", "_rels/.rels"):
                root = etree.fromstring(data)
                for relation in list(root):
                    target = relation.get("Target", "")
                    if "customXml/" in target or target.endswith("docProps/custom.xml"):
                        root.remove(relation)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                      standalone=True)
            elif name == "[Content_Types].xml":
                root = etree.fromstring(data)
                for entry in list(root):
                    part_name = entry.get("PartName", "")
                    if part_name.startswith("/customXml/") or \
                            part_name == "/docProps/custom.xml":
                        root.remove(entry)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                      standalone=True)
            result.writestr(info, data)


def build(reference, source, output):
    doc = Document(reference)
    ref_paragraphs = list(doc.paragraphs)
    ref_table = doc.tables[0]._tbl
    patterns = {
        "chapter": ref_paragraphs[0]._p,
        "major": ref_paragraphs[1]._p,
        "body": ref_paragraphs[2]._p,
        "numbered": ref_paragraphs[7]._p,
        "sub": ref_paragraphs[22]._p,
        "caption": ref_paragraphs[46]._p,
    }
    styles = {
        "chapter": (run_props(ref_paragraphs[0], 0), run_props(ref_paragraphs[0], 1)),
        "major": (run_props(ref_paragraphs[1], 1), run_props(ref_paragraphs[1], 0)),
        "body": (run_props(ref_paragraphs[2], 0), run_props(ref_paragraphs[2], 1)),
    }
    body = doc._body._body
    section_properties = body.find(W + "sectPr")
    for element in list(body):
        if element is not section_properties:
            body.remove(element)
    for kind, value in parse_markdown(source):
        if kind == "table":
            element = make_table(ref_table, value, *styles["body"])
        elif kind == "caption":
            element = clean_clone(patterns["caption"])
            append_runs(element, value[0], *styles["major"])
            append_runs(element, value[1], *styles["body"])
            caption_props = child(element, "pPr")
            child(caption_props, "keepNext")
            # This Chapter 1 schedule does not fit below the preceding prose.
            # Start its caption and complete editable table together.
            if value[0] == "ตารางที่ 1.1":
                child(caption_props, "pageBreakBefore")
        else:
            pattern = patterns.get(kind, patterns["body"])
            if kind == "chapter":
                fonts = styles["chapter"]
            elif kind in ("major", "sub"):
                fonts = styles["major"]
            else:
                fonts = styles["body"]
            element = make_paragraph(pattern, value, *fonts,
                                     left=kind == "sub",
                                     keep_next=kind in ("chapter", "major", "sub"))
        body.insert(len(body) - 1, element)
    page_numbers(doc)
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "E-Filing ISO 9001 Chapter 1 Draft"
    props.subject = ""
    props.keywords = ""
    props.comments = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(dir=output.parent) as temp_dir:
        intermediate = Path(temp_dir) / "reference_based.docx"
        doc.save(intermediate)
        remove_unused_metadata(intermediate, output, reference)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("Usage: script REFERENCE.docx SOURCE.md OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]))
